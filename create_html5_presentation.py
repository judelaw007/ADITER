#!/usr/bin/env python3
"""
ADIT HTML5 Presentation Generator
Creates professional HTML5 presentations with Google TTS voice narration
Integrates HTML5_Templates design with interactive features
"""

import os
import sys
import json
import base64
import hashlib
import requests
from pathlib import Path
from datetime import datetime
import docx
import re


class HTML5PresentationBuilder:
    """Build professional presentations with HTML5_Templates design"""

    def __init__(self, google_api_key: str, output_dir: str = "presentations"):
        self.google_api_key = google_api_key
        self.output_dir = Path(output_dir)
        self.tts_url = "https://texttospeech.googleapis.com/v1/text:synthesize"

    def extract_content_from_docx(self, docx_path: str):
        """Extract structured content from DOCX file"""
        doc = docx.Document(docx_path)
        slides = []
        current_slide = None

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Check if it's a heading
            if para.style.name.startswith('Heading'):
                # Save previous slide if exists
                if current_slide and current_slide['content']:
                    slides.append(current_slide)

                # Start new slide
                heading_level = int(para.style.name.replace('Heading ', ''))
                current_slide = {
                    'title': text,
                    'content': [],
                    'level': heading_level
                }
            else:
                # Add content to current slide
                if current_slide:
                    current_slide['content'].append(text)

        # Add the last slide
        if current_slide and current_slide['content']:
            slides.append(current_slide)

        return slides

    def generate_voice_google(self, text: str, output_path: str, voice: str = "en-US-Neural2-D"):
        """Generate voice narration using Google Cloud TTS"""
        # Limit text length for TTS
        if len(text) > 5000:
            text = text[:5000] + "..."

        data = {
            'input': {'text': text},
            'voice': {
                'languageCode': 'en-US',
                'name': voice
            },
            'audioConfig': {
                'audioEncoding': 'MP3',
                'pitch': 0,
                'speakingRate': 1.0
            }
        }

        try:
            response = requests.post(
                self.tts_url,
                params={'key': self.google_api_key},
                json=data,
                timeout=30
            )

            if response.status_code == 200:
                audio_content = response.json()['audioContent']
                audio_data = base64.b64decode(audio_content)

                with open(output_path, 'wb') as f:
                    f.write(audio_data)

                print(f"✓ Generated audio: {Path(output_path).name}")
                return True
            else:
                print(f"✗ TTS Error: {response.status_code} - {response.text}")
                return False

        except Exception as e:
            print(f"✗ TTS Exception: {str(e)}")
            return False

    def create_html5_presentation(self, slides: list, title: str, topic_number: str):
        """Create HTML5 presentation with template design"""

        # Create output directory
        presentation_dir = self.output_dir / f"{topic_number}_{title.replace(' ', '_')}"
        presentation_dir.mkdir(parents=True, exist_ok=True)
        audio_dir = presentation_dir / "audio"
        audio_dir.mkdir(exist_ok=True)

        # Generate audio for each slide
        audio_files = []
        for idx, slide in enumerate(slides):
            # Combine title and content for narration
            narration_text = f"{slide['title']}. " + " ".join(slide['content'][:5])  # Limit content

            audio_file = audio_dir / f"slide_{idx:03d}.mp3"
            if self.generate_voice_google(narration_text, str(audio_file)):
                audio_files.append(f"audio/slide_{idx:03d}.mp3")
            else:
                audio_files.append(None)

        # Generate HTML with HTML5_Templates design
        html_content = self.generate_html_template(slides, title, topic_number, audio_files)

        # Save HTML file
        html_path = presentation_dir / "index.html"
        with open(html_path, 'w', encoding='utf-8') as f:
            f.write(html_content)

        print(f"\n✅ Presentation created: {html_path}")
        return str(html_path)

    def generate_html_template(self, slides: list, title: str, topic_number: str, audio_files: list):
        """Generate HTML with HTML5_Templates design integration"""

        # Build slides HTML
        slides_html = ""
        for idx, slide in enumerate(slides):
            content_html = ""
            for item in slide['content'][:8]:  # Limit items per slide
                content_html += f"<li>{item}</li>\n"

            audio_tag = ""
            if audio_files[idx]:
                audio_tag = f'<audio id="audio-{idx}" src="{audio_files[idx]}"></audio>'

            slides_html += f'''
            <div class="slide" data-slide="{idx}">
                <div class="slide-content">
                    <h2 class="slide-title">{slide['title']}</h2>
                    <ul class="bullet-points">
                        {content_html}
                    </ul>
                </div>
                {audio_tag}
            </div>
            '''

        # Generate complete HTML with HTML5_Templates styling
        html = f'''<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title} - ADIT Fundamental Tax Issues</title>
    <style>
        * {{
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }}

        :root {{
            --primary: #003F87;
            --secondary: #0066CC;
            --accent: #FF6B35;
            --light-bg: #F8F9FA;
            --dark-text: #1A1F36;
            --light-text: #FFFFFF;
            --success: #27AE60;
        }}

        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, 'Helvetica Neue', Arial, sans-serif;
            background: var(--light-bg);
            color: var(--dark-text);
            line-height: 1.6;
            overflow: hidden;
        }}

        .presentation-container {{
            width: 100vw;
            height: 100vh;
            display: flex;
            flex-direction: column;
            background: linear-gradient(135deg, var(--primary) 0%, var(--secondary) 100%);
        }}

        header {{
            background: rgba(0, 0, 0, 0.3);
            color: var(--light-text);
            padding: 20px 40px;
            display: flex;
            justify-content: space-between;
            align-items: center;
        }}

        .header-title {{
            font-size: 24px;
            font-weight: 700;
        }}

        .header-subtitle {{
            font-size: 14px;
            opacity: 0.9;
        }}

        .slide-counter {{
            font-size: 16px;
            font-weight: 600;
        }}

        .slides-wrapper {{
            flex: 1;
            position: relative;
            overflow: hidden;
            padding: 40px;
        }}

        .slide {{
            position: absolute;
            top: 0;
            left: 0;
            width: 100%;
            height: 100%;
            display: flex;
            align-items: center;
            justify-content: center;
            opacity: 0;
            transform: translateX(100%);
            transition: all 0.6s cubic-bezier(0.4, 0, 0.2, 1);
            padding: 40px;
        }}

        .slide.active {{
            opacity: 1;
            transform: translateX(0);
            z-index: 10;
        }}

        .slide.previous {{
            transform: translateX(-100%);
        }}

        .slide-content {{
            background: white;
            border-radius: 16px;
            padding: 50px;
            max-width: 1000px;
            width: 100%;
            box-shadow: 0 20px 60px rgba(0, 0, 0, 0.3);
            max-height: 80vh;
            overflow-y: auto;
        }}

        .slide-title {{
            color: var(--primary);
            font-size: 32px;
            font-weight: 700;
            margin-bottom: 30px;
            border-bottom: 3px solid var(--accent);
            padding-bottom: 15px;
        }}

        .bullet-points {{
            list-style: none;
            padding: 0;
        }}

        .bullet-points li {{
            padding: 12px 0;
            padding-left: 40px;
            position: relative;
            font-size: 16px;
            line-height: 1.8;
            color: #555;
        }}

        .bullet-points li::before {{
            content: "→";
            position: absolute;
            left: 0;
            color: var(--accent);
            font-weight: bold;
            font-size: 20px;
        }}

        .controls {{
            background: rgba(0, 0, 0, 0.3);
            padding: 20px 40px;
            display: flex;
            justify-content: space-between;
            align-items: center;
        }}

        .nav-button {{
            background: linear-gradient(135deg, var(--primary), var(--secondary));
            color: var(--light-text);
            border: none;
            padding: 12px 30px;
            border-radius: 8px;
            font-size: 16px;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.3s ease;
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.2);
        }}

        .nav-button:hover:not(:disabled) {{
            transform: translateY(-2px);
            box-shadow: 0 6px 18px rgba(0, 0, 0, 0.3);
        }}

        .nav-button:disabled {{
            opacity: 0.5;
            cursor: not-allowed;
        }}

        .audio-control {{
            background: var(--accent);
            color: white;
            border: none;
            padding: 12px 25px;
            border-radius: 8px;
            font-size: 14px;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.3s ease;
        }}

        .audio-control:hover {{
            background: #FF8C5A;
            transform: scale(1.05);
        }}

        .audio-control.playing {{
            animation: pulse 1.5s ease-in-out infinite;
        }}

        @keyframes pulse {{
            0%, 100% {{ opacity: 1; }}
            50% {{ opacity: 0.7; }}
        }}

        .progress-bar {{
            position: fixed;
            top: 0;
            left: 0;
            height: 4px;
            background: var(--accent);
            transition: width 0.3s ease;
            z-index: 1000;
        }}

        @media (max-width: 768px) {{
            .slide-content {{
                padding: 30px;
            }}

            .slide-title {{
                font-size: 24px;
            }}

            .bullet-points li {{
                font-size: 14px;
            }}
        }}
    </style>
</head>
<body>
    <div class="progress-bar" id="progressBar"></div>

    <div class="presentation-container">
        <header>
            <div>
                <div class="header-title">ADIT - Fundamental Tax Issues</div>
                <div class="header-subtitle">{title}</div>
            </div>
            <div class="slide-counter">
                <span id="currentSlide">1</span> / <span id="totalSlides">{len(slides)}</span>
            </div>
        </header>

        <div class="slides-wrapper" id="slidesWrapper">
            {slides_html}
        </div>

        <div class="controls">
            <button class="nav-button" id="prevBtn" onclick="previousSlide()">← Previous</button>
            <button class="audio-control" id="audioBtn" onclick="toggleAudio()">🔊 Audio On</button>
            <button class="nav-button" id="nextBtn" onclick="nextSlide()">Next →</button>
        </div>
    </div>

    <script>
        let currentSlide = 0;
        const totalSlides = {len(slides)};
        let audioEnabled = true;
        let currentAudio = null;

        // Initialize first slide
        document.addEventListener('DOMContentLoaded', function() {{
            showSlide(0);
            updateProgress();
        }});

        function showSlide(index) {{
            const slides = document.querySelectorAll('.slide');

            // Remove active and previous classes
            slides.forEach((slide, idx) => {{
                slide.classList.remove('active', 'previous');
                if (idx < index) {{
                    slide.classList.add('previous');
                }}
            }});

            // Set active slide
            slides[index].classList.add('active');
            currentSlide = index;

            // Update UI
            document.getElementById('currentSlide').textContent = index + 1;
            document.getElementById('prevBtn').disabled = index === 0;
            document.getElementById('nextBtn').disabled = index === totalSlides - 1;

            // Play audio
            playSlideAudio(index);
            updateProgress();
        }}

        function playSlideAudio(index) {{
            // Stop current audio
            if (currentAudio) {{
                currentAudio.pause();
                currentAudio.currentTime = 0;
            }}

            if (!audioEnabled) return;

            // Find and play audio for this slide
            const audio = document.getElementById(`audio-${{index}}`);
            if (audio) {{
                currentAudio = audio;
                audio.play().catch(e => console.log('Audio play failed:', e));
                document.getElementById('audioBtn').classList.add('playing');

                audio.onended = () => {{
                    document.getElementById('audioBtn').classList.remove('playing');
                }};
            }}
        }}

        function nextSlide() {{
            if (currentSlide < totalSlides - 1) {{
                showSlide(currentSlide + 1);
            }}
        }}

        function previousSlide() {{
            if (currentSlide > 0) {{
                showSlide(currentSlide - 1);
            }}
        }}

        function toggleAudio() {{
            audioEnabled = !audioEnabled;
            const btn = document.getElementById('audioBtn');

            if (audioEnabled) {{
                btn.textContent = '🔊 Audio On';
                playSlideAudio(currentSlide);
            }} else {{
                btn.textContent = '🔇 Audio Off';
                if (currentAudio) {{
                    currentAudio.pause();
                    currentAudio.currentTime = 0;
                }}
                btn.classList.remove('playing');
            }}
        }}

        function updateProgress() {{
            const progress = ((currentSlide + 1) / totalSlides) * 100;
            document.getElementById('progressBar').style.width = progress + '%';
        }}

        // Keyboard navigation
        document.addEventListener('keydown', function(e) {{
            if (e.key === 'ArrowRight' || e.key === ' ') {{
                e.preventDefault();
                nextSlide();
            }} else if (e.key === 'ArrowLeft') {{
                e.preventDefault();
                previousSlide();
            }} else if (e.key === 'm' || e.key === 'M') {{
                toggleAudio();
            }}
        }});

        // Touch swipe support
        let touchStartX = 0;
        let touchEndX = 0;

        document.addEventListener('touchstart', e => {{
            touchStartX = e.changedTouches[0].screenX;
        }});

        document.addEventListener('touchend', e => {{
            touchEndX = e.changedTouches[0].screenX;
            handleSwipe();
        }});

        function handleSwipe() {{
            if (touchEndX < touchStartX - 50) {{
                nextSlide();
            }}
            if (touchEndX > touchStartX + 50) {{
                previousSlide();
            }}
        }}
    </script>
</body>
</html>'''

        return html


def main():
    """Main execution function"""

    # Check for Google API key
    api_key = os.getenv('GOOGLE_API_KEY')
    if not api_key:
        print("❌ ERROR: GOOGLE_API_KEY environment variable not set!")
        print("\nSet it with:")
        print("  export GOOGLE_API_KEY='your-key-here'")
        sys.exit(1)

    # Find DOCX files
    source_dir = Path("generated_notes/01_Fundamental_tax_issues/quality_checked")
    if not source_dir.exists():
        print(f"❌ ERROR: Directory not found: {source_dir}")
        sys.exit(1)

    docx_files = list(source_dir.glob("*_Enhanced_*.docx"))
    if not docx_files:
        print(f"❌ ERROR: No DOCX files found in {source_dir}")
        sys.exit(1)

    print(f"📚 Found {len(docx_files)} DOCX files")
    print("=" * 60)

    # Create output directory
    output_dir = Path("generated_notes/01_Fundamental_tax_issues/presentations")
    output_dir.mkdir(parents=True, exist_ok=True)

    # Initialize builder
    builder = HTML5PresentationBuilder(api_key, str(output_dir))

    # Process each DOCX file
    for idx, docx_path in enumerate(sorted(docx_files), 1):
        print(f"\n📄 [{idx}/{len(docx_files)}] Processing: {docx_path.name}")
        print("-" * 60)

        # Extract topic info from filename
        filename = docx_path.stem
        parts = filename.replace("Fundamental_tax_issues_", "").replace("_Enhanced_2025-11-14", "")
        topic_title = parts.replace("_", " ")
        topic_number = f"{idx:02d}"

        print(f"📌 Topic: {topic_title}")

        # Extract content
        print("📖 Extracting content from DOCX...")
        slides = builder.extract_content_from_docx(str(docx_path))
        print(f"✓ Extracted {len(slides)} slides")

        # Create presentation
        print(f"🎨 Creating HTML5 presentation...")
        presentation_path = builder.create_html5_presentation(
            slides=slides,
            title=topic_title,
            topic_number=topic_number
        )

        print(f"✅ Presentation complete!")
        print(f"📍 Location: {presentation_path}")

    print("\n" + "=" * 60)
    print("🎉 ALL PRESENTATIONS CREATED SUCCESSFULLY!")
    print("=" * 60)
    print(f"\n📂 Output directory: {output_dir}")
    print("\n🌐 To view presentations:")
    print(f"   cd {output_dir}/[topic_folder]")
    print("   python3 -m http.server 8000")
    print("   Open: http://localhost:8000")


if __name__ == "__main__":
    main()
